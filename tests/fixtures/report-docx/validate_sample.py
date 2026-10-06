"""Read-only interoperability checks for the opt-in Word export sample."""
from pathlib import Path
import json
import zipfile
import xml.etree.ElementTree as ET
from docx import Document

folder = Path(__file__).parent
sample = folder / 'sample.docx'
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(sample) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        ET.fromstring(archive.read(name))
    root = ET.fromstring(archive.read('word/document.xml'))
    paragraphs = [''.join(p.itertext()) for p in root.findall('.//w:body/w:p', ns)]
    text = '\n'.join(paragraphs)
    assert len(archive.namelist()) == 8
    assert 'SUCCESS_SENTINEL' not in text
    assert all('原编号 ' + n in text for n in ['0012', '0042', '0219', '1024'])
    assert '<A&B>' in text
    assert '尝试后失败 1 条；尚未处理 1 条；缺少下载入口 1 条；处理中 1 条' in text
    links = ET.fromstring(archive.read('word/_rels/document.xml.rels'))
    external = {r.attrib['Id']: r.attrib['Target'] for r in links if r.attrib.get('TargetMode') == 'External'}
    assert any('ABCdef%2F123&source=author' in value for value in external.values())
    for h in root.findall('.//w:hyperlink', ns):
        assert h.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'] in external

document = Document(sample)
assert len(document.tables) == 0
assert document.paragraphs[0].text == '未获取全文文献清单'
assert document.core_properties.title == '未获取全文文献清单'
assert round(document.sections[0].page_width.mm) == 210
assert round(document.sections[0].page_height.mm) == 297
result = {'zip_crc': 'passed', 'all_xml': 'parsed', 'python_docx': 'opened',
          'paragraphs': len(document.paragraphs), 'unresolved_records': 4,
          'external_links': len(external), 'page_mm': [210, 297], 'tables': 0}
(folder / 'validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(result, ensure_ascii=False))
